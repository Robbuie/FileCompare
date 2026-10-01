"""Excel workbooks for table compare (1.3), through openpyxl.

A workbook pair opens in the same key-matched grid as two CSV files: one
sheet at a time, chosen by name, its first used row as the column names,
each record found by its key wherever it is. What a cell holds is compared
as text the way a person reads it -- `12` not `12.0`, a date as
`2026-09-30`, TRUE and FALSE -- so a number typed as 12 and one calculated
as 12.0 are the same, and "Numbers by value" still applies on top.

**Values, or formulas.** By default a cell is the value Excel last
calculated and saved with the file (`data_only`), which is what somebody
looking at the sheet sees. "Formulas" compares what was typed instead, so a
changed formula that happens to give the same answer still shows.

Only `.xlsx`, `.xlsm`, `.xltx` and `.xltm`. The old binary `.xls` needs a
different reader and is said so rather than guessed at.

Pure apart from openpyxl: bytes in, `tables.Table`s out. Runs in the loader.
"""

from __future__ import annotations

import datetime as _dt
import io
from dataclasses import dataclass
from functools import lru_cache

from app.core.tables import Table

EXTENSIONS = frozenset({"xlsx", "xlsm", "xltx", "xltm"})
OLD = frozenset({"xls", "xlt"})

#: Rows and columns past which a sheet is cut short, and says so. A sheet
#: with a million rows is a database, and the grid is not one.
MAX_ROWS = 200_000
MAX_COLUMNS = 500


def _ext(path: str) -> str:
    name = path.replace("\\", "/").rsplit("/", 1)[-1]
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def is_workbook(path: str) -> bool:
    return _ext(path) in EXTENSIONS


def is_old_workbook(path: str) -> bool:
    return _ext(path) in OLD


def text(value) -> str:
    """A cell as a person reads it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, _dt.timedelta):
        return str(value)
    return str(value)


@dataclass
class Book:
    """Every sheet of one workbook, read once."""

    names: list[str]
    sheets: dict[str, Table]
    #: Per sheet, why it was cut short, when it was.
    cut: dict[str, str]


def read(data: bytes, *, formulas: bool = False) -> Book:
    """All sheets. Raises ValueError with a sentence when it is not a
    workbook openpyxl can read."""
    return _read(data, formulas)


@lru_cache(maxsize=4)
def _read(data: bytes, formulas: bool) -> Book:
    # Keyed on the bytes themselves: two reads of the same file -- a key
    # changed, a toggle flipped -- cost one parse. Four entries is two
    # workbooks in each of two modes.
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - shipped in the installer
        raise ValueError("Excel workbooks need openpyxl, which is not installed") from exc
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                      data_only=not formulas, keep_links=False)
    except Exception as exc:  # noqa: BLE001 - zip, XML and openpyxl's own errors
        raise ValueError(f"Not a workbook that can be read: {exc}") from None
    names: list[str] = []
    sheets: dict[str, Table] = {}
    cut: dict[str, str] = {}
    try:
        for sheet in book.worksheets:
            names.append(sheet.title)
            rows: list[list[str]] = []
            widest = 0
            for number, values in enumerate(sheet.iter_rows(values_only=True)):
                if number >= MAX_ROWS:
                    cut[sheet.title] = f"only the first {MAX_ROWS:,} rows are compared"
                    break
                cells = [text(v) for v in values[:MAX_COLUMNS]]
                if len(values) > MAX_COLUMNS:
                    cut[sheet.title] = f"only the first {MAX_COLUMNS} columns are compared"
                while cells and cells[-1] == "":
                    cells.pop()
                rows.append(cells)
                widest = max(widest, len(cells))
            while rows and not rows[-1]:
                rows.pop()
            # The first row with anything in it is where the table starts:
            # a title row or two of blank lines above a schedule is common.
            while rows and not rows[0]:
                rows.pop(0)
            rows = [row + [""] * (widest - len(row)) for row in rows]
            sheets[sheet.title] = Table(header=rows[0] if rows else [],
                                        rows=rows[1:] if rows else [])
    finally:
        book.close()
    return Book(names, sheets, cut)


def table(book: Book, name: str, *, header: bool = True) -> Table:
    """One sheet as a table; a sheet this book does not have is empty. With
    `header` off the first row is data, and the columns are numbered."""
    found = book.sheets.get(name)
    if found is None:
        return Table(header=[], rows=[])
    if header:
        return found
    rows = ([found.header] if found.header else []) + found.rows
    width = max((len(r) for r in rows), default=0)
    return Table(header=[f"Column {i + 1}" for i in range(width)], rows=rows)


@dataclass
class SheetState:
    name: str
    left: bool
    right: bool
    same: bool | None       # None when only one side has it


def sheet_states(left: Book, right: Book) -> list[SheetState]:
    """Every sheet name, left order first, and whether its cells match."""
    order = list(left.names) + [n for n in right.names if n not in left.names]
    out = []
    for name in order:
        a, b = left.sheets.get(name), right.sheets.get(name)
        same = None
        if a is not None and b is not None:
            same = a.header == b.header and a.rows == b.rows
        out.append(SheetState(name, a is not None, b is not None, same))
    return out


def first_different(states: list[SheetState]) -> str:
    """The sheet to open on: the first that differs, else the first."""
    for state in states:
        if state.same is False or state.same is None:
            return state.name
    return states[0].name if states else ""
